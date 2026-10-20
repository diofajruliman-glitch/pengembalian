"""Create an explicitly marked working draft, not an approved letter template."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent.parent
report=Document();section=report.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(2.4);section.bottom_margin=Cm(2.4);section.left_margin=Cm(2.8);section.right_margin=Cm(2.2)
style=report.styles['Normal'];style.font.name='Arial';style.font.size=Pt(11);style.paragraph_format.space_after=Pt(6)
style.paragraph_format.line_spacing=1.15
for text in ['DRAF — BELUM DISAHKAN','LAPORAN HASIL PEMERIKSAAN','DIREKTORAT PERLINDUNGAN SOSIAL NON KEBENCANAAN']:
 p=report.add_paragraph();p.alignment=1;r=p.add_run(text);r.bold=True
p=report.add_paragraph('Nomor: {nomor}');p.alignment=1
report.add_paragraph('Dokumen kerja berdasarkan data wawancara. Memerlukan pemeriksaan dan pengesahan pejabat berwenang sebelum digunakan sebagai LHP resmi.')
for heading,key in [('I. DASAR DAN RUANG LINGKUP PEMERIKSAAN','dasar'),('II. IDENTITAS ASN','identitas'),('III. TEMUAN AWAL','temuan'),('IV. FAKTA HASIL WAWANCARA','fakta'),('V. ANALISIS DAN PERTIMBANGAN','analisis'),('VI. KESIMPULAN DAN REKOMENDASI','kesimpulan')]:
 p=report.add_paragraph();p.add_run(heading).bold=True
 report.add_paragraph('{'+key+'}')
 if key=='kesimpulan':
  report.add_paragraph('Rekomendasi tindak lanjut:').runs[0].bold=True
  report.add_paragraph('{rekomendasi}')
report.add_paragraph('Status dokumen pendukung: {status_sptjm}')
report.add_paragraph('Referensi sumber: {referensi}')
p=report.add_paragraph('{kota}, {tanggal}');p.alignment=2
p=report.add_paragraph('Pemeriksa / Penelaah,\n\n\n(....................................)\nNIP. ................................');p.alignment=2
footer=section.footer.paragraphs[0];footer.alignment=1
footer.add_run('DRAF — BELUM DISAHKAN · Halaman ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
report.save(root/'public/templates/lhp-draft-2026.docx')
